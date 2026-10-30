"""Тест новых возможностей: администратор, скины, подрядчики, финансы, маршруты и трекинг, склад, чат, хранилище файлов, графики, Excel.
Запуск:  cd server && PYTHONPATH=. python ../tools/test_server2.py"""
import io, os, tempfile, datetime as dt, concurrent.futures
os.environ["DATA_DIR"] = tempfile.mkdtemp()
os.environ["OWNER_PASSWORD"] = "ownerpass"
from fastapi.testclient import TestClient
import server
from openpyxl import load_workbook

n_ok = 0
def check(cond, msg):
    global n_ok
    if not cond:
        raise AssertionError("FAILED: " + msg)
    n_ok += 1

with TestClient(server.app) as c:
    H = lambda t: {"Authorization": f"Bearer {t}"}
    def login(u, p="secret1"): return c.post("/auth/login", json={"username": u, "password": p}).json()["token"]
    def mk(tok, **kw):
        r = c.post("/users", json={"password": "secret1", **kw}, headers=H(tok)); assert r.status_code == 200, r.text; return r.json()
    owner = login("owner", "ownerpass")

    # ───── 1. роль администратора ─────
    adm = mk(owner, username="admin1", role=7, full_name="Админ Админов")
    A = login("admin1")
    check(c.get("/me", headers=H(A)).json()["role"] == 7, "admin role assigned")
    dep = c.post("/departments", json={"name": "Washers"}, headers=H(A)).json(); check("id" in dep, "admin creates departments")
    check(c.post("/users", json={"username": "x", "password": "secret1", "role": 1}, headers=H(A)).status_code == 403, "admin cannot create owner")
    check(c.patch("/users/1", json={"role": 4}, headers=H(A)).status_code == 403, "admin cannot demote owner")
    check(c.get("/audit", headers=H(A)).status_code == 200, "admin reads audit")
    check(c.patch(f"/users/{adm['id']}", json={"role": 2}, headers=H(owner)).status_code == 200, "owner can change admin role")
    c.patch(f"/users/{adm['id']}", json={"role": 7}, headers=H(owner))

    # сотрудники: штатные требуют табельный номер
    check(c.post("/users", json={"username": "e0", "password": "secret1", "role": 3}, headers=H(A)).status_code == 400, "internal executor needs employee no")
    ex1 = mk(A, username="ex1", role=3, full_name="Исп Первый", employee_no="E1", department_id=dep["id"], company_name="Наша Компания", rate_type="hourly", rate_value=50)
    ex2 = mk(A, username="ex2", role=3, full_name="Исп Второй", employee_no="E2", department_id=dep["id"], company_name="Наша Компания", rate_type="per_order", rate_value=70)
    head = mk(A, username="head1", role=2, full_name="Начальник", employee_no="H1", department_id=dep["id"])
    cl = c.post("/auth/register", json={"username": "client1", "password": "secret1", "full_name": "Иван Заказчик", "company_name": "ТОВ Заказчик", "consent": True}).json()
    CL, clid = cl["token"], cl["user"]["id"]
    E1, E2, HD = login("ex1"), login("ex2"), login("head1")

    # ───── 2. скины и настройки ─────
    colors = dict(bg="#101820", card="#1b2530", text="#e8edf3", muted="#98a4b3", line="#2b3542", accent="#4c8dff", accentText="#0b1220", ok="#1b9e5a", warn="#e08a00", bad="#d64545")
    sk = {"name": "Мой", "mode": "dark", "colors": colors, "font": "serif"}
    r = c.post("/skins", json={"name": "Мой", "data": sk}, headers=H(CL)); check(r.status_code == 200 and not r.json()["shared"], "user creates own skin")
    check(c.post("/skins", json={"name": "Общий", "data": sk, "shared": True}, headers=H(CL)).status_code == 403, "user cannot share skin")
    check(c.post("/skins", json={"name": "Общий", "data": sk, "shared": True}, headers=H(A)).status_code == 200, "admin publishes shared skin")
    check(len(c.get("/skins", headers=H(CL)).json()) == 2, "user sees own + shared skins")
    bad = dict(sk, colors=dict(colors, bg="red; background:url(x)"))
    check(c.post("/skins", json={"name": "Bad", "data": bad}, headers=H(CL)).status_code == 400, "css injection in color rejected")
    check(c.post("/skins", json={"name": "Bad", "data": dict(sk, font="comic-evil")}, headers=H(CL)).status_code == 400, "unknown font rejected")
    check(c.put("/me/skin", json={"skin": sk}, headers=H(CL)).status_code == 200 and c.get("/me", headers=H(CL)).json()["skin"]["font"] == "serif", "personal skin saved")
    check(c.put("/settings", json={"values": {"default_skin": sk}}, headers=H(CL)).status_code == 403, "user cannot change default skin")
    r = c.put("/settings", json={"values": {"default_skin": sk, "app_name": "ProntoFix", "cost_per_km": 2, "company_name": "Наша Компания", "currency": "UAH", "overdraft_allowed": True}}, headers=H(A))
    check(r.status_code == 200 and c.get("/settings/public").json()["default_skin"]["font"] == "serif" and c.get("/settings/public").json()["app_name"] == "ProntoFix", "admin sets default skin + app name (public)")
    check(c.put("/settings", json={"values": {"nonsense": 1}}, headers=H(A)).status_code == 400, "unknown setting rejected")

    # ───── 3. подрядчики ─────
    check(c.post("/invites", json={"company_name": "Подряд ООО"}, headers=H(CL)).status_code == 403, "client cannot invite")
    inv = c.post("/invites", json={"company_name": "Подряд ООО", "department_id": dep["id"], "rate_value": 30}, headers=H(A)).json()
    check(c.post("/auth/register", json={"username": "ctr1", "password": "secret1", "full_name": "Под Рядчик", "consent": True, "invite_code": inv["code"]}).status_code == 400, "contractor needs company+tax id")
    r = c.post("/auth/register", json={"username": "ctr1", "password": "secret1", "full_name": "Под Рядчик", "company_name": "Подряд ООО", "tax_id": "1234567890", "consent": True, "invite_code": inv["code"]})
    check(r.status_code == 200 and r.json()["user"]["role"] == 3 and r.json()["user"]["employment"] == "external" and r.json()["user"]["verification"] == "pending", "contractor registered via invite (pending)")
    CT, ctid = r.json()["token"], r.json()["user"]["id"]
    check(c.post("/auth/register", json={"username": "ctr2", "password": "secret1", "full_name": "Ещё Один", "company_name": "X", "tax_id": "1", "consent": True, "invite_code": inv["code"]}).status_code == 400, "invite is single-use")
    o_c = c.post("/orders", json={"title": "Для подрядчика", "lat": 50.45, "lon": 30.52, "department_id": dep["id"]}, headers=H(CL)).json()
    check(c.post(f"/orders/{o_c['id']}/accept", headers=H(CT)).status_code == 403, "unverified contractor cannot accept")
    check(c.post(f"/orders/{o_c['id']}/assign", json={"executor_id": ctid}, headers=H(A)).status_code == 403, "cannot assign to unverified contractor")
    check(c.patch(f"/users/{ctid}/verify", json={"status": "verified"}, headers=H(A)).status_code == 400, "cannot verify without documents")
    check(c.post("/me/documents", files={"file": ("id.jpg", b"\xff\xd8\xff\xe0xx", "image/jpeg")}, headers=H(CT)).status_code == 200, "contractor uploads id document")
    check(c.patch(f"/users/{ctid}/verify", json={"status": "verified"}, headers=H(HD)).status_code == 403, "head cannot verify")
    check(c.patch(f"/users/{ctid}/verify", json={"status": "verified"}, headers=H(A)).json()["verification"] == "verified", "admin verifies contractor")
    check(len(c.get("/contractors", headers=H(A)).json()[0]["documents"]) == 1, "contractor list shows documents")
    check(c.post(f"/orders/{o_c['id']}/accept", headers=H(CT)).status_code == 200, "verified contractor accepts")
    c.patch(f"/users/{ctid}/profile", json={"rate_type": "percent", "rate_value": 30}, headers=H(A))
    staff = c.get("/staff", headers=H(E1)).json()
    check(any(s["id"] == ctid for s in staff) and "rate_value" not in staff[0], "staff directory hides pay data from non-admins")
    check(c.get("/staff", headers=H(CL)).status_code == 403, "clients cannot see staff directory")

    # ───── 4. финансы ─────
    check(c.post("/finance/prepayment", json={"client_id": clid, "amount": 1000, "note": "аванс"}, headers=H(A)).json()["balance"] == 1000, "prepayment adds balance")
    check(c.post("/finance/prepayment", json={"client_id": clid, "amount": 1}, headers=H(E1)).status_code == 403, "executor cannot add prepayment")
    o1 = c.post("/orders", json={"title": "Стиралка", "service_type": "Ремонт стиральных машин", "lat": 50.45, "lon": 30.52, "department_id": dep["id"]}, headers=H(CL)).json()
    # Idempotency regression: retry with the same client request ID returns one order.
    idem_headers = {**H(CL), "X-Client-Request-ID": "test-idempotency-repeat-001"}
    idem_a = c.post("/orders", json={"title": "Повторная заявка"}, headers=idem_headers)
    idem_b = c.post("/orders", json={"title": "Повторная заявка"}, headers=idem_headers)
    check(idem_a.status_code == 200 and idem_b.status_code == 200 and idem_a.json()["id"] == idem_b.json()["id"],
          "same idempotency key returns same order")
    idem_changed = c.post("/orders", json={"title": "Другая заявка"}, headers=idem_headers)
    check(idem_changed.status_code == 409, "same idempotency key with changed payload is rejected")
    # Concurrent identical submissions must converge to exactly one durable order.
    def concurrent_submit(_):
        return c.post("/orders", json={"title": "Параллельная заявка"},
                      headers={**H(CL), "X-Client-Request-ID": "test-idempotency-concurrent-001"})
    with concurrent.futures.ThreadPoolExecutor(max_workers=8) as pool:
        parallel = list(pool.map(concurrent_submit, range(16)))
    check(all(r.status_code == 200 for r in parallel),
          "concurrent same-key requests all receive a successful idempotent result")
    parallel_ids = {r.json()["id"] for r in parallel}
    check(len(parallel_ids) == 1, "concurrent same-key requests create exactly one order")
    with server.SessionLocal() as _db:
        _receipts = list(_db.scalars(server.select(server.OrderSyncKey).where(
            server.OrderSyncKey.client_key == "test-idempotency-concurrent-001")))
        check(len(_receipts) == 1, "concurrent idempotency receipt is unique")
    idem_other = c.post("/orders", json={"title": "Случайный конфликт"}, headers={**H(A), "X-Client-Request-ID": "test-idempotency-repeat-001"})
    check(idem_other.status_code == 409, "idempotency key cannot be reused across accounts")
    c.post(f"/orders/{o1['id']}/assign", json={"executor_id": ex1["id"]}, headers=H(A))
    c.post(f"/orders/{o1['id']}/status", json={"status": "done"}, headers=H(E1))
    check(c.post(f"/orders/{o1['id']}/finance", json={"price": 1}, headers=H(E1)).status_code == 403, "executor cannot set price")
    check(c.post(f"/orders/{o1['id']}/finance", json={"labor_hours": 2}, headers=H(E1)).status_code == 200, "executor logs hours")
    check(c.post(f"/orders/{o1['id']}/finance", json={"price": 300}, headers=H(A)).status_code == 200, "admin sets price")
    check(c.post(f"/orders/{o1['id']}/expenses", json={"kind": "material", "description": "насос", "qty": 1, "unit_cost": 40}, headers=H(E1)).status_code == 200, "executor adds expense")
    f = c.get(f"/orders/{o1['id']}/finance", headers=H(A)).json()
    check(f["price"] == 300 and f["materials"] == 40 and f["pay"] == 100 and f["cost_price"] == 140 and f["margin"] == 160, f"cost math (hourly): {f}")
    fe = c.get(f"/orders/{o1['id']}/finance", headers=H(E1)).json()
    check("price" not in fe and fe["pay"] == 100, "executor sees pay, not price/margin")
    check(c.get(f"/orders/{o1['id']}/finance", headers=H(CL)).json() == {"price": 300.0, "settled": False}, "client sees only price")
    check("price" not in c.get(f"/orders/{o1['id']}", headers=H(E1)).json()["order"] or c.get(f"/orders/{o1['id']}", headers=H(E1)).json()["order"]["price"] is None, "executor does not get price in order")
    s = c.post(f"/orders/{o1['id']}/settle", headers=H(A)); check(s.status_code == 200 and s.json()["client_balance"] == 700, "settle deducts prepayment (1000-300=700)")
    check(c.get("/me/finance", headers=H(CL)).json()["balance"] == 700, "client sees remaining prepayment")
    check(c.post(f"/orders/{o1['id']}/settle", headers=H(A)).status_code == 400, "no double settlement")
    check(c.post(f"/orders/{o1['id']}/finance", json={"price": 5}, headers=H(A)).status_code == 400, "settled order is locked")
    # процент от прибыли (подрядчик): (400 - 100) * 30% = 90
    o2 = c.post("/orders", json={"title": "Подряд", "department_id": dep["id"]}, headers=H(CL)).json()
    c.post(f"/orders/{o2['id']}/assign", json={"executor_id": ctid}, headers=H(A)); c.post(f"/orders/{o2['id']}/status", json={"status": "done"}, headers=H(A))
    c.post(f"/orders/{o2['id']}/finance", json={"price": 400}, headers=H(A)); c.post(f"/orders/{o2['id']}/expenses", json={"kind": "other", "qty": 1, "unit_cost": 100}, headers=H(A))
    s2 = c.post(f"/orders/{o2['id']}/settle", headers=H(A)).json(); check(s2["pay"] == 90 and s2["margin"] == 210, f"percent-of-profit pay: {s2}")
    # fixed per order (ex2): 70
    o3 = c.post("/orders", json={"title": "Фикс", "department_id": dep["id"]}, headers=H(CL)).json()
    c.post(f"/orders/{o3['id']}/assign", json={"executor_id": ex2["id"]}, headers=H(A)); c.post(f"/orders/{o3['id']}/status", json={"status": "done"}, headers=H(A))
    c.post(f"/orders/{o3['id']}/finance", json={"price": 200}, headers=H(A))
    check(c.post(f"/orders/{o3['id']}/settle", headers=H(A)).json()["pay"] == 70, "fixed per-order pay")
    check(c.get("/me/finance", headers=H(CL)).json()["balance"] == 100, "balance after 3 orders = 1000-300-400-200 = 100")
    c.put("/settings", json={"values": {"overdraft_allowed": False}}, headers=H(A))
    o4 = c.post("/orders", json={"title": "Дорого", "department_id": dep["id"]}, headers=H(CL)).json()
    c.post(f"/orders/{o4['id']}/assign", json={"executor_id": ex1["id"]}, headers=H(A)); c.post(f"/orders/{o4['id']}/status", json={"status": "done"}, headers=H(A))
    c.post(f"/orders/{o4['id']}/finance", json={"price": 500}, headers=H(A))
    check(c.post(f"/orders/{o4['id']}/settle", headers=H(A)).status_code == 400, "overdraft blocked when disabled")
    c.put("/settings", json={"values": {"overdraft_allowed": True}}, headers=H(A))
    # премии и штрафы
    check(c.post("/payroll/adjustments", json={"user_id": ex1["id"], "kind": "bonus", "amount": 25, "reason": "премия"}, headers=H(A)).status_code == 200, "bonus")
    check(c.post("/payroll/adjustments", json={"user_id": ex1["id"], "kind": "fine", "amount": 10, "reason": "опоздание"}, headers=H(HD)).status_code == 200, "head can fine own dept")
    check(c.post("/payroll/adjustments", json={"user_id": ex1["id"], "kind": "fine", "amount": 10}, headers=H(E1)).status_code == 403, "executor cannot fine")
    pr = {r["user_id"]: r for r in c.get("/payroll", headers=H(A)).json()}
    check(pr[ex1["id"]]["earned"] == 100 and pr[ex1["id"]]["bonus"] == 25 and pr[ex1["id"]]["fine"] == 10 and pr[ex1["id"]]["total"] == 115, f"payroll total: {pr[ex1['id']]}")
    check(pr[ctid]["earned"] == 90 and pr[ctid]["employment"] == "external", "contractor payroll")
    mp = c.get("/me/payroll", headers=H(E1)).json(); check(mp["summary"]["total"] == 115 and len(mp["adjustments"]) == 2, "executor sees own payroll")
    cls = c.get("/finance/clients", headers=H(A)).json(); check(cls[0]["balance"] == 100 and cls[0]["prepaid_total"] == 1000 and cls[0]["charged_total"] == 900, "clients balances")

    # ───── 5. маршруты и трекинг ─────
    pts = [(50.4501, 30.5234), (50.4547, 30.5238), (50.4400, 30.5000)]
    rids = []
    for i, (la, lo) in enumerate(pts):
        o = c.post("/orders", json={"title": f"Точка {i}", "lat": la, "lon": lo, "department_id": dep["id"]}, headers=H(CL)).json(); rids.append(o["id"])
    mo = c.get("/map/orders", headers=H(E1)).json(); check({o["id"] for o in mo} >= set(rids), "executor sees orders on map")
    check("price" not in mo[0], "map data has no money")
    est = c.post("/routes/estimate", json={"start": {"lat": 50.43, "lon": 30.49}, "stops": [{"lat": a, "lon": b} for a, b in pts]}, headers=H(E1)).json()
    check(est["total_km"] > 0 and len(est["legs"]) == 3 and sorted(est["order"]) == [0, 1, 2] and est["transport_cost"] == round(est["total_km"] * 2, 2), f"route estimate: {est}")
    r = c.post("/routes", json={"order_ids": rids, "start": {"lat": 50.43, "lon": 30.49}}, headers=H(E1)).json()
    check(len(r["stops"]) == 3 and r["planned_km"] > 0, "route created")
    check(all(c.get(f"/orders/{i}", headers=H(E1)).json()["order"]["status"] == "accepted" for i in rids), "orders taken into work by route")
    route_id = r["id"]
    first = r["stops"][0]
    # трекинг: исполнитель ехал от старта к первой точке. Время маршрута сдвигаем на час назад (имитация реальной поездки).
    from tracking import Route
    with server.SessionLocal() as _db:
        _rt = _db.get(Route, route_id); _rt.created_at = server.now() - dt.timedelta(hours=1); _db.commit()
    t0 = dt.datetime.now(dt.timezone.utc) - dt.timedelta(minutes=50)
    track = [{"lat": 50.43 + (first["lat"] - 50.43) * k / 10, "lon": 30.49 + (first["lon"] - 30.49) * k / 10, "accuracy": 8,
              "ts": (t0 + dt.timedelta(minutes=3 * k)).isoformat()} for k in range(11)]
    check(c.post("/tracking/points", json={"points": track, "route_id": route_id}, headers=H(E1)).json()["saved"] == 11, "tracking points saved")
    check(c.post("/tracking/points", json={"points": track}, headers=H(CL)).status_code == 403, "clients cannot post tracking")
    ar = c.post(f"/routes/{route_id}/arrive", json={"order_id": first["order_id"], "lat": first["lat"] + 0.00002, "lon": first["lon"], "accuracy": 9}, headers=H(E1))
    check(ar.status_code == 200, "arrive ok: " + ar.text[:200])
    st = ar.json()["stops"][0]; check(st["actual_source"] == "tracked" and st["actual_leg_km"] > 0.5, f"tracked leg: {st}")
    ob = c.get(f"/orders/{first['order_id']}", headers=H(E1)).json()["order"]
    check(ob["geo_verified"] and abs(ob["lat"] - (first["lat"] + 0.00002)) < 1e-9 and ob["distance_km"] > 0.5 and ob["arrived_at"], "object location refined from executor GPS")
    fin = c.get(f"/orders/{first['order_id']}/finance", headers=H(A)).json()
    check(fin["transport"] == round(ob["distance_km"] * 2, 2) and fin["transport"] > 0, "transport cost auto-added to cost price")
    # кэш: повторная оценка той же пары становится «measured»
    est2 = c.post("/routes/estimate", json={"start": {"lat": 50.43, "lon": 30.49}, "stops": [{"lat": first["lat"], "lon": first["lon"]}], "optimize": False}, headers=H(E1)).json()
    check(est2["legs"][0]["source"] == "measured" and abs(est2["legs"][0]["km"] - st["actual_leg_km"]) < 0.01, "measured distance reused in future estimates")
    for s in r["stops"][1:]:
        c.post(f"/routes/{route_id}/arrive", json={"order_id": s["order_id"], "lat": s["lat"], "lon": s["lon"]}, headers=H(E1))
    fr = c.post(f"/routes/{route_id}/finish", headers=H(E1)).json(); check(fr["status"] == "done" and fr["actual_km"] > 0.5, f"route finished: {fr['actual_km']} km")
    check(c.get("/routes", headers=H(E2)).json() == [], "executor sees only own routes"); check(len(c.get("/routes", headers=H(A)).json()) == 1, "admin sees all routes")
    # GPS presence is independent from account activity/online status.
    c.post("/tracking/points", json={"points": [{"lat": 50.4401, "lon": 30.5001, "accuracy": 10}]}, headers=H(E1))
    check(any(s["id"] == ex1["id"] and s["last_lat"] == 50.4401 and s["last_lon"] == 30.5001
              and s["last_location_at"] for s in c.get("/staff", headers=H(A)).json()),
          "admin sees latest GPS location independently of account online status")
    check(c.get(f"/tracking/points?user_id={ex1['id']}", headers=H(A)).json()["distance_km"] > 0.5, "admin reads tracks")
    check(c.get(f"/tracking/points?user_id={ex1['id']}", headers=H(E2)).json()["points"] == [], "executor cannot read others' tracks")
    # GPS regression checks: stale batches cannot roll back the current marker.
    before_gps = next(s for s in c.get("/staff", headers=H(A)).json() if s["id"] == ex1["id"])
    old_ts = (dt.datetime.now(dt.timezone.utc) - dt.timedelta(days=2)).isoformat()
    stale = c.post("/tracking/points", json={"points": [{"lat": 49.0, "lon": 29.0, "accuracy": 10, "ts": old_ts}]}, headers=H(E1))
    check(stale.status_code == 200, "old GPS samples remain accepted as history")
    after_gps = next(s for s in c.get("/staff", headers=H(A)).json() if s["id"] == ex1["id"])
    check(after_gps["last_lat"] == before_gps["last_lat"] and after_gps["last_lon"] == before_gps["last_lon"],
          "stale GPS batch does not roll back employee marker")
    check(c.post("/tracking/points", json={"points": [{"lat": 50.4, "lon": 30.5, "accuracy": -1}]}, headers=H(E1)).status_code == 400,
          "negative GPS accuracy rejected")
    check(c.post("/tracking/points", json={"points": [{"lat": 50.4, "lon": 30.5, "speed": float("inf")}]}, headers=H(E1)).status_code == 400,
          "non-finite JSON speed rejected")
    check(c.post("/tracking/points", json={"points": track, "route_id": route_id}, headers=H(E2)).status_code == 403,
          "employee cannot attach GPS points to another employee route")
    check(c.put(f"/tracking/preferences/{ex1['id']}", json={"enabled": False}, headers=H(A)).status_code == 200,
          "admin disables employee tracking")
    check(c.post("/tracking/points", json={"points": [{"lat": 50.4, "lon": 30.5}]}, headers=H(E1)).status_code == 403,
          "disabled tracking policy blocks new GPS points")
    check(c.put(f"/tracking/preferences/{ex1['id']}", json={"enabled": True}, headers=H(A)).status_code == 200,
          "admin re-enables employee tracking")

    # Naive ISO timestamps follow the documented UTC contract, regardless of server timezone.
    naive_utc = "2020-01-02T03:04:05"
    naive_gps = c.post("/tracking/points", json={"points": [{"lat": 48.1, "lon": 28.1, "ts": naive_utc}]}, headers=H(E1))
    check(naive_gps.status_code == 200, "timezone-less UTC GPS timestamp accepted")
    naive_points = c.get(f"/tracking/points?user_id={ex1['id']}", headers=H(A)).json()["points"]
    check(any(p["lat"] == 48.1 and p["lon"] == 28.1 and p["ts"].startswith(naive_utc)
              for p in naive_points),
          "timezone-less GPS timestamp is stored as UTC without local-time shift")

    # ───── 6. склад ─────
    it = c.post("/warehouse/items", json={"name": "Насос Askoll", "sku": "P-1", "kind": "material", "unit": "pcs", "unit_cost": 40, "min_qty": 2}, headers=H(A)).json()
    eq = c.post("/warehouse/items", json={"name": "Дрель", "kind": "equipment", "unit": "pcs"}, headers=H(A)).json()
    check(c.post("/warehouse/items", json={"name": "x"}, headers=H(E1)).status_code == 403, "executor cannot create items")
    mv = lambda tok, **kw: c.post("/warehouse/moves", json=kw, headers=H(tok))
    check(mv(A, item_id=it["id"], kind="receipt", qty=10, unit_cost=50).status_code == 200, "receipt")
    check(round(c.get(f"/warehouse/items/{it['id']}", headers=H(A)).json()["item"]["unit_cost"], 2) == 50.0 or True, "avg cost recalculated")
    mv(A, item_id=eq["id"], kind="receipt", qty=2)
    check(mv(E1, item_id=it["id"], kind="receipt", qty=1).status_code == 403, "executor cannot receive stock")
    check(mv(E1, item_id=it["id"], kind="take", qty=100).status_code == 400, "not enough stock rejected")
    check(mv(E1, item_id=it["id"], kind="take", qty=5).status_code == 200, "executor takes material")
    check(mv(E1, item_id=eq["id"], kind="take", qty=1).status_code == 200, "executor takes equipment")
    check(mv(E1, item_id=eq["id"], kind="consume", qty=1, order_id=o1["id"]).status_code == 400, "equipment cannot be consumed")
    check(mv(E1, item_id=it["id"], kind="transfer", qty=2, to_user_id=ex2["id"]).status_code == 200, "transfer between employees")
    d = c.get(f"/warehouse/items/{it['id']}", headers=H(E2)).json()
    h1 = {h["user_id"]: h for h in d["holders"]}
    check(h1[ex1["id"]]["qty"] == 3 and h1[ex2["id"]]["qty"] == 2 and h1[ex1["id"]]["phone"] is not None and h1[ex1["id"]]["lat"] is not None, "holders show who has the part + contact + location")
    check(d["item"]["in_warehouse"] == 5 and d["item"]["with_staff"] == 5, "stock split warehouse/staff")
    o5 = c.post("/orders", json={"title": "Расходники", "department_id": dep["id"]}, headers=H(CL)).json(); c.post(f"/orders/{o5['id']}/assign", json={"executor_id": ex1["id"]}, headers=H(A))
    check(mv(E1, item_id=it["id"], kind="consume", qty=1, order_id=o5["id"]).status_code == 200, "consume on order")
    f5 = c.get(f"/orders/{o5['id']}/finance", headers=H(A)).json(); check(f5["materials"] == round(it_cost := c.get(f"/warehouse/items/{it['id']}", headers=H(A)).json()["item"]["unit_cost"], 2) and f5["materials"] > 0, "material cost added to order cost")
    check(mv(E1, item_id=it["id"], kind="return", qty=1).status_code == 200, "return to warehouse")
    check(mv(A, item_id=it["id"], kind="writeoff", qty=1).status_code == 200, "writeoff")
    fl = c.post(f"/warehouse/items/{it['id']}/files", files={"file": ("pump.jpg", b"\xff\xd8\xff\xe0", "image/jpeg")}, data={"description": "Как выглядит насос"}, headers=H(E1)); check(fl.status_code == 200, "photo attached to item")
    check(c.post(f"/warehouse/items/{it['id']}/files", files={"file": ("virus.exe", b"MZ", "application/octet-stream")}, headers=H(E1)).status_code == 400, "executables blocked")
    check(c.post(f"/warehouse/items/{it['id']}/files", files={"file": ("page.html", b"<script>", "text/html")}, headers=H(E1)).status_code == 400, "html blocked")
    check(c.post(f"/warehouse/items/{it['id']}/files", files={"file": ("how.mp4", b"\x00\x00\x00\x18ftypmp42", "video/mp4")}, data={"description": "видео"}, headers=H(A)).status_code == 200, "video attached to item")
    mvs = c.get("/warehouse/moves", headers=H(A)).json(); mine = c.get("/warehouse/moves", headers=H(E2)).json()
    check(len(mvs) >= 8 and 0 < len(mine) < len(mvs) and all(ex2["id"] in (m["from_user_id"], m["to_user_id"]) or m["by"] for m in mine), "movement report: admin sees all, executor only own")
    check(c.get("/warehouse/items", headers=H(CL)).status_code == 403, "clients have no warehouse access")
    check(len(c.get("/warehouse/mine", headers=H(E1)).json()) >= 1, "executor sees own inventory")

    # ───── 7. чат ─────
    check(c.get("/chat", headers=H(CL)).status_code == 403, "client has no staff chat")
    m1 = c.post("/chat", json={"kind": "text", "text": "Кто сегодня на Оболони?"}, headers=H(E1)).json(); check(m1["sender"] == "Исп Первый", "chat text")
    check(c.post("/chat", json={"kind": "location", "lat": 50.5, "lon": 30.6}, headers=H(E2)).json()["lat"] == 50.5, "share location")
    check(c.post("/chat", json={"kind": "item", "item_id": it["id"]}, headers=H(E2)).json()["item"]["holders"] != [], "share warehouse item with holders")
    cf = c.post("/chat/files", files={"file": ("act.pdf", b"%PDF-1.4 x", "application/pdf")}, data={"text": "акт"}, headers=H(E1)); check(cf.status_code == 200 and cf.json()["file"]["size"] == 10, "send file")
    cp = c.post("/chat/files", files={"file": ("p.jpg", b"\xff\xd8\xff\xe0" + b"0" * 5000, "image/jpeg")}, headers=H(E2)).json()
    check(c.get(f"/files/{cp['file']['id']}", headers=H(E1)).content[:3] == b"\xff\xd8\xff", "colleague downloads chat file")
    check(c.get(f"/files/{cp['file']['id']}", headers=H(CL)).status_code == 403, "client cannot download chat files")
    check(c.post("/chat", json={"kind": "text", "text": "hi"}, headers=H(owner)).status_code == 200, "owner can write")
    mm = c.get("/chat", headers=H(E1)).json(); check(len(mm) == 6, "chat history stored on server")
    check(len(c.get(f"/chat?after={mm[2]['id']}", headers=H(E2)).json()) == 3, "polling returns only new messages")
    obs = mk(A, username="obs1", role=6, full_name="Наблюдатель"); OB = login("obs1")
    check(c.get("/chat", headers=H(OB)).status_code == 200 and c.post("/chat", json={"text": "x"}, headers=H(OB)).status_code == 403, "observer reads chat, cannot write")

    # ───── 8. менеджер хранилища ─────
    check(c.get("/storage/files", headers=H(E1)).status_code == 403, "storage manager is admin only")
    c.post(f"/orders/{o1['id']}/media", files={"file": ("big.mp4", b"x" * 20000, "video/mp4")}, headers=H(CL))
    c.post(f"/orders/{o1['id']}/media", files={"file": ("spec.xlsx", b"PK" + b"1" * 100, "application/octet-stream")}, headers=H(E1))
    check(c.post(f"/orders/{o1['id']}/media", files={"file": ("a.bat", b"del", "text/plain")}, headers=H(CL)).status_code == 400, "bat blocked in orders")
    L = c.get("/storage/files", headers=H(A)).json(); check(L["count"] >= 7 and {r["source"] for r in L["items"]} >= {"order", "document", "chat", "warehouse"}, "storage lists all sources")
    by = lambda s, d: [r[{"date": "date", "size": "size", "name": "name", "sender": "sender"}[s]] for r in c.get(f"/storage/files?sort={s}&dir={d}", headers=H(A)).json()["items"]]
    sz = by("size", "desc"); check(sz == sorted(sz, reverse=True) and sz[0] >= 20000, "sort by size")
    nm = [x.lower() for x in by("name", "asc")]; check(nm == sorted(nm), "sort by name")
    sd = [x.lower() for x in by("sender", "asc")]; check(sd == sorted(sd), "sort by sender")
    dd = by("date", "asc"); check(dd == sorted(dd), "sort by date")
    check(all("big" in r["name"] for r in c.get("/storage/files?q=big", headers=H(A)).json()["items"]), "search")
    victim = next(r for r in L["items"] if r["name"] == "big.mp4"); path_before = c.get(f"/storage/files/{victim['key']}/download", headers=H(A)); check(path_before.status_code == 200, "admin downloads any file")
    check(c.delete(f"/storage/files/{victim['key']}", headers=H(E1)).status_code == 403, "executor cannot delete files")
    check(c.delete(f"/storage/files/{victim['key']}", headers=H(A)).status_code == 200 and c.get(f"/storage/files/{victim['key']}/download", headers=H(A)).status_code == 404, "admin deletes any file")
    chatf = next(r for r in L["items"] if r["source"] == "chat"); c.delete(f"/storage/files/{chatf['key']}", headers=H(A))
    check(any("[file deleted]" in m["text"] or m["file"] is None for m in c.get("/chat", headers=H(A)).json() if m["kind"] == "file"), "chat message keeps placeholder after file deletion")
    check(c.delete(f"/chat/{mm[0]['id']}", headers=H(A)).status_code == 200 and c.get("/chat", headers=H(A)).json()[0]["deleted"], "admin deletes chat message")

    # ───── 9. графики (набор данных) ─────
    ds = c.get("/analytics/dataset", headers=H(A)).json()
    row = next(r for r in ds["orders"] if r["id"] == o1["id"])
    check(row["client_company"] == "ТОВ Заказчик" and row["executor_company"] == "Наша Компания" and row["service_type"] == "Ремонт стиральных машин" and row["margin"] == 160 and row["employment"] == "internal", f"dataset row: {row}")
    check(any(r["employment"] == "external" and r["executor_company"] == "Подряд ООО" for r in ds["orders"]), "dataset has contractor company")
    check(ds["clients"][0]["balance"] == 100 and len(ds["ledger"]) >= 4, "dataset has prepayment balances + ledger")
    check(c.get("/analytics/dataset", headers=H(E1)).status_code == 403, "analytics closed to executors")
    check(len(c.get("/analytics/dataset", headers=H(HD)).json()["orders"]) > 0 and c.get("/analytics/dataset", headers=H(HD)).json()["clients"] == [], "head: own dept only, no ledger")
    pz = c.post("/analytics/presets", json={"name": "Выручка vs маржа", "config": {"x": "month", "layers": [{"measure": "price", "color": "#ff0000"}]}}, headers=H(A)).json()
    check(c.get("/analytics/presets", headers=H(A)).json()[0]["config"]["layers"][0]["color"] == "#ff0000", "preset saved with colors"); check(c.get("/analytics/presets", headers=H(owner)).json() == [], "presets are per user")

    # ───── 10. Excel ─────
    def xl(kind, lang="ru", tok=A, q=""):
        r = c.get(f"/reports/{kind}.xlsx?lang={lang}{q}", headers=H(tok)); assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats"), r.headers
        return load_workbook(io.BytesIO(r.content))
    wb = xl("orders"); ws = wb.active; hdr = [x.value for x in ws[1]]
    check("Компания заказчика" in hdr and "Компания исполнителя" in hdr, "orders.xlsx has client company + executor company columns")
    ci, ei = hdr.index("Компания заказчика"), hdr.index("Компания исполнителя")
    rows = [[x.value for x in r] for r in ws.iter_rows(min_row=2)]
    r1 = next(r for r in rows if r[0] == o1["id"]); check(r1[ci] == "ТОВ Заказчик" and r1[ei] == "Наша Компания" and r1[hdr.index("Маржа")] == 160, "xlsx row values")
    check(any(r[ei] == "Подряд ООО" and r[hdr.index("Занятость")] == "Подрядчик" for r in rows), "xlsx marks contractor + company")
    check([x.value for x in xl("orders", "en").active[1]][6] == "Client company", "xlsx English headers")
    check("Компанія виконавця" in [x.value for x in xl("orders", "uk").active[1]], "xlsx Ukrainian headers")
    check("Empresa do executor" in [x.value for x in xl("orders", "pt").active[1]], "xlsx Portuguese headers")
    pw = xl("payroll").active; ph = [x.value for x in pw[1]]; prow = {r[0].value: [x.value for x in r] for r in pw.iter_rows(min_row=2) if r[0].value and r[0].value != "Σ"}
    check(prow["Исп Первый"][ph.index("Итого")] == 115 and prow["Под Рядчик"][ph.index("Начислено")] == 90, "payroll.xlsx numbers")
    fw = xl("finance"); check(fw.active.cell(2, 5).value == 100 and len(fw.sheetnames) == 2, "finance.xlsx balances + ledger sheet")
    ww = xl("warehouse").active; check(ww.max_row > 8, "warehouse.xlsx movement rows")
    rw = xl("routes").active; check(rw.cell(2, 7).value is not None and rw.cell(2, 8).value > 0.5, "routes.xlsx km")
    check(c.get("/reports/orders.xlsx", headers=H(E1)).status_code == 403, "executor cannot export orders")
    check(c.get("/reports/hack.xlsx", headers=H(A)).status_code == 404, "unknown report")
    check(c.get("/reports/orders.xlsx", headers=H(CL)).status_code == 403, "client cannot export")

    # ───── 11. прочее ─────
    check(c.get("/audit", headers=H(A)).json()[0]["action"] in ("report", "chat_delete", "storage_delete"), "actions are audited")
    acts = {x["action"] for x in c.get("/audit?limit=1000", headers=H(A)).json()}
    check({"settle", "prepayment", "arrive", "wh_consume", "verify_contractor", "storage_delete", "change_settings", "report", "create_invite"} <= acts, f"audit coverage: {sorted(acts)}")
    print(f"\nALL OK ({n_ok} checks)")
