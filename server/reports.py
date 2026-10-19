"""
Excel-отчёты (.xlsx). Заголовки колонок — на языке интерфейса (параметр lang: ru | uk | en | pt).

  orders       все заказы: заказчик + КОМПАНИЯ ЗАКАЗЧИКА, исполнитель + КОМПАНИЯ ИСПОЛНИТЕЛЯ, занятость, деньги, километры
  payroll      оплата сотрудников: штат (ставка) и подрядчики (% от прибыли), премии и штрафы
  finance      остатки предоплаты заказчиков и история проводок
  warehouse    движение материалов и оборудования
  routes       маршруты и километраж исполнителей
Файл отдаётся приложением; куда его сохранить (папка/носитель), пользователь выбирает на своём устройстве.
"""
import datetime as dt
import io
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from server import ADMINS, CLIENT, HEAD, OBSERVER, STAFF_ROLES, EXECUTOR, User, current_user, get_db, log, need

router = APIRouter()

_L = """no|№|№|No.|Nº
date|Дата|Дата|Date|Data
status|Статус|Статус|Status|Status
service|Вид услуги|Вид послуги|Service type|Tipo de serviço
department|Направление|Напрям|Department|Departamento
client|Заказчик|Замовник|Client|Cliente
client_company|Компания заказчика|Компанія замовника|Client company|Empresa do cliente
address|Адрес|Адреса|Address|Endereço
executor|Исполнитель|Виконавець|Executor|Executor
executor_company|Компания исполнителя|Компанія виконавця|Executor company|Empresa do executor
employment|Занятость|Зайнятість|Employment|Vínculo
hours|Часы|Години|Hours|Horas
price|Цена услуги|Ціна послуги|Service price|Preço do serviço
materials|Материалы|Матеріали|Materials|Materiais
transport|Транспорт|Транспорт|Transport|Transporte
other|Прочее|Інше|Other|Outros
pay|Оплата исполнителю|Оплата виконавцю|Executor pay|Pagamento ao executor
cost_price|Себестоимость|Собівартість|Cost price|Custo
margin|Маржа|Маржа|Margin|Margem
km|Км|Км|Km|Km
rating|Оценка|Оцінка|Rating|Avaliação
settled|Рассчитан|Розраховано|Settled|Liquidado
employee|Сотрудник|Співробітник|Employee|Funcionário
company|Компания|Компанія|Company|Empresa
rate_type|Тип ставки|Тип ставки|Rate type|Tipo de taxa
rate|Ставка|Ставка|Rate|Taxa
orders_count|Заказов|Замовлень|Orders|Pedidos
earned|Начислено|Нараховано|Earned|Ganho
bonus|Премии|Премії|Bonuses|Bônus
fine|Штрафы|Штрафи|Fines|Multas
total|Итого|Разом|Total|Total
balance|Остаток предоплаты|Залишок передоплати|Prepayment balance|Saldo de pré-pagamento
prepaid|Внесено предоплаты|Внесено передоплати|Prepaid total|Total pré-pago
charged|Списано за услуги|Списано за послуги|Charged for services|Cobrado
kind|Тип|Тип|Type|Tipo
amount|Сумма|Сума|Amount|Valor
note|Примечание|Примітка|Note|Observação
order|Заказ|Замовлення|Order|Pedido
item|Позиция|Позиція|Item|Item
sku|Артикул|Артикул|SKU|SKU
qty|Количество|Кількість|Quantity|Quantidade
unit|Ед.|Од.|Unit|Un.
from|Откуда|Звідки|From|De
to|Куда|Куди|To|Para
by|Кто провёл|Хто провів|Done by|Feito por
route|Маршрут|Маршрут|Route|Rota
planned_km|План, км|План, км|Planned km|Km planejado
actual_km|Факт, км|Факт, км|Actual km|Km real
stops|Заездов|Заїздів|Stops|Paradas
yes|да|так|yes|sim
no_|нет|ні|no|não
internal|Штат|Штат|Staff|Equipe própria
external|Подрядчик|Підрядник|Contractor|Terceirizado
hourly|Почасовая|Погодинна|Hourly|Por hora
per_order|За заказ|За замовлення|Per order|Por pedido
percent|% от прибыли|% від прибутку|% of profit|% do lucro
st_new|Новый|Новий|New|Novo
st_assigned|Назначен|Призначено|Assigned|Atribuído
st_accepted|Принят|Прийнято|Accepted|Aceito
st_in_progress|В работе|У роботі|In progress|Em andamento
st_estimate|Оценка объёма|Оцінка обсягу|Estimate|Orçamento
st_done|Выполнен|Виконано|Done|Concluído
st_cancelled|Отменён|Скасовано|Cancelled|Cancelado
receipt|Поступление|Надходження|Receipt|Entrada
take|Выдача|Видача|Issue|Retirada
transfer|Передача|Передача|Transfer|Transferência
return|Возврат|Повернення|Return|Devolução
consume|Списание на заказ|Списання на замовлення|Used on order|Usado no pedido
writeoff|Списание|Списання|Write-off|Baixa
prepayment|Предоплата|Передоплата|Prepayment|Pré-pagamento
charge|Списание за услугу|Списання за послугу|Service charge|Cobrança
adjust|Корректировка|Коригування|Adjustment|Ajuste
active|В пути|У дорозі|Active|Ativa
done|Завершён|Завершено|Finished|Finalizada
sheet_orders|Заказы|Замовлення|Orders|Pedidos
sheet_payroll|Оплата сотрудников|Оплата співробітників|Payroll|Folha
sheet_finance|Предоплата|Передоплата|Prepayment|Pré-pagamento
sheet_ledger|Проводки|Проводки|Ledger|Lançamentos
sheet_warehouse|Склад|Склад|Warehouse|Armazém
sheet_routes|Маршруты|Маршрути|Routes|Rotas
"""
LABELS = {"ru": {}, "uk": {}, "en": {}, "pt": {}}
for _line in _L.strip().splitlines():
    _k, *_vals = _line.split("|")
    for _lang, _v in zip(("ru", "uk", "en", "pt"), _vals):
        LABELS[_lang][_k] = _v

HEADER_FILL = PatternFill("solid", fgColor="2F6FED")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="DDE2EA")
MONEY = '#,##0.00'


def tr(lang: str):
    d = LABELS.get(lang) or LABELS["en"]
    return lambda k: d.get(k, LABELS["en"].get(k, k))


def build_sheet(ws, headers, rows, money_cols=(), number_cols=(), total_cols=(), widths=None):
    """Заголовок, строки, формат чисел, закрепление шапки, фильтр и строка итогов."""
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for r in rows:
        ws.append(r)
    n = len(rows)
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, max_row=n + 1, min_col=col, max_col=col):
            row[0].number_format = MONEY
    for col in number_cols:
        for row in ws.iter_rows(min_row=2, max_row=n + 1, min_col=col, max_col=col):
            row[0].number_format = '0.00'
    if total_cols and n:
        ws.append([None] * len(headers))
        tot = n + 2
        ws.cell(tot, 1, "Σ").font = Font(bold=True)
        for col in total_cols:
            L = get_column_letter(col)
            cell = ws.cell(tot, col, f"=SUM({L}2:{L}{n + 1})")
            cell.font = Font(bold=True)
            cell.number_format = MONEY if col in money_cols else '0.00'
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(i)
        if not w:
            w = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 9), 48)
    ws.freeze_panes = "A2"
    if n:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n + 1}"


def stream(wb: Workbook, name: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fn = f"{name}_{dt.date.today().isoformat()}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fn}"'})


@router.get("/reports/{kind}.xlsx")
def report(kind: str, lang: str = "ru", frm: Optional[str] = None, to: Optional[str] = None,
           user: User = Depends(current_user), db: Session = Depends(get_db)):
    t = tr(lang)
    wb = Workbook()
    ws = wb.active
    if kind == "orders":
        need(user, *ADMINS, OBSERVER, HEAD)
        from analytics import order_rows
        data = order_rows(db, user, frm, to)
        ws.title = t("sheet_orders")
        headers = [t(k) for k in ("no", "date", "status", "service", "department", "client", "client_company", "address", "executor",
                                  "executor_company", "employment", "hours", "price", "materials", "transport", "other", "pay",
                                  "cost_price", "margin", "km", "rating", "settled")]
        rows = [[r["id"], r["date"], t("st_" + r["status"]), r["service_type"], r["department"], r["client_name"], r["client_company"],
                 r["address"], r["executor"], r["executor_company"], t(r["employment"]) if r["employment"] in ("internal", "external") else r["employment"],
                 r["hours"], r["price"], r["materials"], r["transport"], r["other"], r["pay"], r["cost"], r["margin"], r["km"],
                 r["rating"], t("yes") if r["settled"] else t("no_")] for r in data]
        build_sheet(ws, headers, rows, money_cols=(13, 14, 15, 16, 17, 18, 19), number_cols=(12, 20), total_cols=(12, 13, 14, 15, 16, 17, 18, 19, 20),
                    widths={7: 28, 10: 28, 8: 30})
        name = "orders"
    elif kind == "payroll":
        need(user, *ADMINS, OBSERVER)
        from finance import payroll_rows
        data = payroll_rows(db, frm, to)
        ws.title = t("sheet_payroll")
        headers = [t(k) for k in ("employee", "company", "employment", "rate_type", "rate", "orders_count", "earned", "bonus", "fine", "total")]
        rows = [[r["name"], r["company"], t(r["employment"]), t(r["rate_type"]), r["rate_value"], r["orders"], r["earned"], r["bonus"], r["fine"], r["total"]] for r in data]
        build_sheet(ws, headers, rows, money_cols=(7, 8, 9, 10), total_cols=(6, 7, 8, 9, 10), widths={1: 28, 2: 28})
        name = "payroll"
    elif kind == "finance":
        need(user, *ADMINS, OBSERVER)
        from finance import Transaction, balance_of, r2
        from sqlalchemy import func
        ws.title = t("sheet_finance")
        rows = []
        for c in db.scalars(select(User).where(User.role == CLIENT).order_by(User.company_name, User.full_name)):
            prepaid = db.scalar(select(func.coalesce(func.sum(Transaction.amount), 0.0)).where(Transaction.client_id == c.id, Transaction.kind == "prepayment"))
            charged = db.scalar(select(func.coalesce(-func.sum(Transaction.amount), 0.0)).where(Transaction.client_id == c.id, Transaction.kind == "charge"))
            rows.append([c.company_name or "", c.full_name or c.username, r2(prepaid), r2(charged), balance_of(db, c.id)])
        build_sheet(ws, [t("company"), t("client"), t("prepaid"), t("charged"), t("balance")], rows, money_cols=(3, 4, 5), total_cols=(3, 4, 5), widths={1: 30, 2: 28})
        ws2 = wb.create_sheet(t("sheet_ledger"))
        names = {c.id: (c.company_name or c.full_name or c.username) for c in db.scalars(select(User).where(User.role == CLIENT))}
        q = select(Transaction).order_by(Transaction.id)
        from analytics import parse_day
        a, b = parse_day(frm), parse_day(to, True)
        if a:
            q = q.where(Transaction.created_at >= a)
        if b:
            q = q.where(Transaction.created_at < b)
        lrows = [[x.created_at.strftime("%Y-%m-%d %H:%M"), names.get(x.client_id, ""), t(x.kind), r2(x.amount), x.order_id, x.note] for x in db.scalars(q.limit(50000))]
        build_sheet(ws2, [t("date"), t("company"), t("kind"), t("amount"), t("order"), t("note")], lrows, money_cols=(4,), total_cols=(4,), widths={2: 30, 6: 40})
        name = "finance"
    elif kind == "warehouse":
        need(user, *STAFF_ROLES)
        from warehouse import Item, Move, move_out
        ws.title = t("sheet_warehouse")
        q = select(Move).order_by(Move.id)
        from analytics import parse_day
        a, b = parse_day(frm), parse_day(to, True)
        if a:
            q = q.where(Move.created_at >= a)
        if b:
            q = q.where(Move.created_at < b)
        if user.role == EXECUTOR:
            q = q.where((Move.from_user_id == user.id) | (Move.to_user_id == user.id) | (Move.created_by == user.id))
        items = {i.id: i for i in db.scalars(select(Item))}
        rows = []
        for m in db.scalars(q.limit(50000)):
            o = move_out(db, m)
            it = items.get(m.item_id)
            rows.append([o["created_at"][:16].replace("T", " "), o["item"], it.sku if it else "", t(m.kind), m.qty, o["unit"],
                         o["from_name"] or "", o["to_name"] or "", m.order_id, m.unit_cost, o["by"], m.note])
        build_sheet(ws, [t(k) for k in ("date", "item", "sku", "kind", "qty", "unit", "from", "to", "order", "price", "by", "note")], rows,
                    money_cols=(10,), widths={2: 30, 7: 22, 8: 22, 12: 36})
        name = "warehouse"
    elif kind == "routes":
        need(user, *ADMINS, OBSERVER)
        from tracking import Route
        import json
        ws.title = t("sheet_routes")
        users = {u.id: u for u in db.scalars(select(User))}
        rows = []
        for r in db.scalars(select(Route).order_by(Route.id)):
            u = users.get(r.user_id)
            rows.append([r.id, (u.full_name or u.username) if u else "", (u.company_name or "") if u else "", r.created_at.strftime("%Y-%m-%d"),
                         t(r.status), len(json.loads(r.stops_json)), r.planned_km, r.actual_km])
        build_sheet(ws, [t(k) for k in ("route", "executor", "executor_company", "date", "status", "stops", "planned_km", "actual_km")], rows,
                    number_cols=(7, 8), total_cols=(6, 7, 8), widths={2: 28, 3: 28})
        name = "routes"
    else:
        raise HTTPException(404, "Unknown report")
    log(db, user, "report", "report", None, kind)
    db.commit()
    return stream(wb, name)
